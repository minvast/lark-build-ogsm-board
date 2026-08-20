#!/usr/bin/env python3
"""Normalize a local OGSM Excel/CSV source and emit a precheck report."""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any


CANONICAL_FIELDS = [
    "department", "person", "period", "objective", "goal", "strategy", "measure",
    "action", "owner", "collaborators", "start_date", "due_date", "status", "progress",
    "risk", "next_action", "evidence", "strategy_source_text",
]

DISPLAY_NAMES = {
    "department": "部门", "person": "个人", "period": "周期", "objective": "O｜目的",
    "goal": "G｜目标", "strategy": "S｜策略", "measure": "M｜衡量", "action": "行动",
    "owner": "负责人原文", "collaborators": "协作人原文", "start_date": "开始日期",
    "due_date": "截止日期", "status": "进度状态", "progress": "进度原值",
    "risk": "风险/卡点", "next_action": "下一步动作", "evidence": "交付物/证据",
    "strategy_source_text": "S｜策略原文（备查）",
}

ALIASES = {
    "department": ["部门", "所属部门", "责任部门", "组织", "department"],
    "person": ["个人", "姓名", "员工", "人员", "person"],
    "period": ["周期", "年度", "时间周期", "ogsm周期", "period"],
    "objective": ["o", "o目的", "o经营目的", "o经营命题", "目的", "经营目的", "objective"],
    "goal": ["g", "g目标", "g主题", "目标", "goal", "goals"],
    "strategy": ["s", "s策略", "s主题", "策略", "strategy", "strategies"],
    "measure": ["m", "m衡量", "m指标", "衡量", "衡量指标", "指标", "指标名称", "关键结果", "measure", "measures"],
    "action": ["行动", "行动项", "行动名称", "任务", "任务名称", "工作项", "工作项名称", "待办事项", "具体行动", "执行动作", "action"],
    "owner": ["负责人", "主负责人", "责任人", "匹配负责人", "确认负责人", "owner"],
    "collaborators": ["协作人", "配合人", "参与人", "collaborators"],
    "start_date": ["开始日期", "计划开始", "开始时间", "startdate"],
    "due_date": ["截止日期", "固定截止日期", "确认截止日期", "建议截止", "完成日期", "截止时间", "计划完成时间", "deadline", "duedate"],
    "status": ["进度状态", "状态", "任务状态", "status"],
    "progress": ["进度", "当前进度", "完成度", "progress"],
    "risk": ["风险卡点", "风险", "卡点", "阻塞点", "问题", "risk"],
    "next_action": ["下一步动作", "下一步", "后续动作", "nextaction"],
    "evidence": ["交付物证据", "成果证据", "交付物", "证据", "链接", "evidence"],
}

HIERARCHY_FIELDS = ["department", "person", "period", "objective", "goal", "strategy", "measure"]
SUPPORTED_STATUS = {"待更新", "未开始", "已启动", "进行中", "接近完成", "待验收", "受阻", "已完成"}


def clean_header(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.strip().lower().replace("｜", "").replace("|", "")
    return re.sub(r"[\s_\-—：:（）()\[\]/\\]+", "", text)


ALIAS_LOOKUP = {clean_header(alias): field for field, aliases in ALIASES.items() for alias in aliases}


def display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_date(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, dt.date):
        return value.isoformat()
    text = display(value)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日", "%m/%d/%Y"):
        try:
            return dt.datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return text


def normalize_progress(value: Any) -> Any:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return round(number * 100, 2) if 0 <= number <= 1 else round(number, 2)
    text = display(value).replace("％", "%")
    match = re.fullmatch(r"\s*(-?\d+(?:\.\d+)?)\s*%?\s*", text)
    return float(match.group(1)) if match else text


def rows_from_csv(path: Path) -> list[tuple[str, list[list[Any]]]]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [(path.stem, [row for row in csv.reader(handle, delimiter=delimiter)])]


def rows_from_xlsx(path: Path) -> list[tuple[str, list[list[Any]]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Reading .xlsx requires openpyxl in the selected Python runtime") from exc
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheets: list[tuple[str, list[list[Any]]]] = []
    for sheet in workbook.worksheets:
        values = [[sheet.cell(row=r, column=c).value for c in range(1, sheet.max_column + 1)] for r in range(1, sheet.max_row + 1)]
        for merged in sheet.merged_cells.ranges:
            top = sheet.cell(merged.min_row, merged.min_col).value
            for r in range(merged.min_row, merged.max_row + 1):
                for c in range(merged.min_col, merged.max_col + 1):
                    values[r - 1][c - 1] = top
        sheets.append((sheet.title, values))
    return sheets


def detect_header(rows: list[list[Any]], forced_row: int | None) -> tuple[int, dict[int, str]] | None:
    indexes = [forced_row - 1] if forced_row else range(min(40, len(rows)))
    best: tuple[int, int, dict[int, str]] | None = None
    for idx in indexes:
        if idx < 0 or idx >= len(rows):
            continue
        mapping: dict[int, str] = {}
        for col, value in enumerate(rows[idx]):
            field = ALIAS_LOOKUP.get(clean_header(value))
            if field and field not in mapping.values():
                mapping[col] = field
        score = len(mapping) + (3 if "action" in mapping.values() else 0)
        if "action" in mapping.values() and len(mapping) >= 3 and (best is None or score > best[0]):
            best = (score, idx, mapping)
    return (best[1], best[2]) if best else None


def stable_key(record: dict[str, Any]) -> str:
    parts = [record.get(field, "") for field in ("department", "period", "objective", "goal", "strategy", "measure", "action")]
    raw = "\x1f".join(str(p).strip().lower() for p in parts)
    return "OGSM-" + hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16].upper()


def column_label(index: int) -> str:
    """Convert a zero-based column index to an Excel-style label."""
    value = index + 1
    label = ""
    while value:
        value, remainder = divmod(value - 1, 26)
        label = chr(65 + remainder) + label
    return label


def probable_multi_action(text: str) -> bool:
    if not text:
        return False
    return bool(re.search(r"(?:^|\n)\s*(?:\d+[、.)）]|[-•])\s*\S+", text)) and "\n" in text


def parse_source(path: Path, forced_header: int | None) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[str]]:
    if path.suffix.lower() in {".csv", ".tsv"}:
        sheets = rows_from_csv(path)
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        sheets = rows_from_xlsx(path)
    else:
        raise ValueError("Supported source types: .xlsx, .xlsm, .csv, .tsv")

    records: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    skipped_sheets: list[str] = []
    for sheet_name, rows in sheets:
        detected = detect_header(rows, forced_header)
        if not detected:
            skipped_sheets.append(sheet_name)
            continue
        header_idx, mapping = detected
        inherited = {field: "" for field in HIERARCHY_FIELDS}
        for row_idx, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
            values = {field: display(row[col]) if col < len(row) else "" for col, field in mapping.items()}
            for field in HIERARCHY_FIELDS:
                if values.get(field):
                    inherited[field] = values[field]
                elif field in mapping.values():
                    values[field] = inherited[field]
            action = values.get("action", "").strip()
            if not action:
                continue
            record = {field: values.get(field, "") for field in CANONICAL_FIELDS}
            record["start_date"] = normalize_date(record["start_date"])
            record["due_date"] = normalize_date(record["due_date"])
            record["progress"] = normalize_progress(record["progress"])
            record["strategy_source_text"] = record["strategy"]
            mapped_columns = sorted(mapping)
            source_range = f"{column_label(mapped_columns[0])}{row_idx}:{column_label(mapped_columns[-1])}{row_idx}"
            record.update({"source_file": path.name, "source_sheet": sheet_name, "source_range": source_range, "source_row": row_idx})
            record["source_record_id"] = stable_key(record)
            records.append(record)

    if not records:
        issues.append({"severity": "blocking", "code": "no_action_table", "message": "No sheet with a recognizable action header and OGSM columns was found."})
    return records, issues, skipped_sheets


def validate(records: list[dict[str, Any]], issues: list[dict[str, Any]], scope: str) -> None:
    required_review = ["objective", "strategy", "measure", "owner", "due_date"]
    for record in records:
        location = f"{record['source_sheet']}!row {record['source_row']}"
        for field in required_review:
            if not record.get(field):
                issues.append({"severity": "review", "code": f"missing_{field}", "source_record_id": record["source_record_id"], "message": f"{location}: {DISPLAY_NAMES[field]} is blank."})
        if scope == "department" and not record.get("department"):
            issues.append({"severity": "review", "code": "missing_department", "source_record_id": record["source_record_id"], "message": f"{location}: 部门 is blank."})
        if record.get("status") and record["status"] not in SUPPORTED_STATUS:
            issues.append({"severity": "review", "code": "unsupported_status", "source_record_id": record["source_record_id"], "message": f"{location}: status '{record['status']}' needs mapping."})
        if probable_multi_action(record.get("action", "")):
            issues.append({"severity": "review", "code": "possible_multi_action", "source_record_id": record["source_record_id"], "message": f"{location}: the action cell may contain multiple actions and should be reviewed before splitting."})
        owner = record.get("owner", "")
        if re.search(r"[?？]|\binhouse\b\s*\d*\s*名|待定|未定|待确认", owner, re.IGNORECASE):
            issues.append({"severity": "review", "code": "unresolved_owner_placeholder", "source_record_id": record["source_record_id"], "message": f"{location}: owner '{owner}' is a role/headcount placeholder; confirm leaving the people field blank."})
        elif re.search(r"[、,，;/；和&＋+]", owner):
            issues.append({"severity": "review", "code": "multiple_primary_owners", "source_record_id": record["source_record_id"], "message": f"{location}: multiple names appear in the primary owner field; choose one owner and move others to collaborators."})

    by_key: dict[str, list[dict[str, Any]]] = defaultdict(list)
    by_strategy: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        by_key[record["source_record_id"]].append(record)
        if record.get("strategy"):
            by_strategy[record["strategy"]].append(record)
    for key, group in by_key.items():
        if len(group) > 1:
            locations = ", ".join(f"{r['source_sheet']}!{r['source_row']}" for r in group)
            issues.append({"severity": "blocking", "code": "duplicate_source_record_id", "source_record_id": key, "message": f"Duplicate action identity at {locations}. Confirm whether these are duplicates or distinct actions."})
    for strategy, group in by_strategy.items():
        measures = {r.get("measure", "") for r in group if r.get("measure")}
        if len(group) >= 10 or len(measures) >= 4:
            issues.append({"severity": "review", "code": "broad_strategy", "message": f"Strategy '{strategy}' contains {len(group)} actions across {len(measures)} measures; review whether child strategies are needed."})
    if scope == "department":
        departments = sorted({record.get("department", "").strip() for record in records if record.get("department", "").strip()})
        if len(departments) > 1:
            issues.append({"severity": "review", "code": "multiple_departments_topology", "message": f"Source contains {len(departments)} departments ({', '.join(departments)}); confirm separate Bases or one combined Base."})


def write_outputs(path: Path, output_dir: Path, scope: str, records: list[dict[str, Any]], issues: list[dict[str, Any]], skipped: list[str]) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    counts = Counter(issue["severity"] for issue in issues)
    fingerprint = hashlib.sha256(path.read_bytes()).hexdigest()
    payload = {
        "schema_version": 1,
        "source": {"file": path.name, "sha256": fingerprint, "scope": scope},
        "summary": {
            "record_count": len(records),
            "strategy_count": len({r["strategy"] for r in records if r.get("strategy")}),
            "owner_text_count": len({r["owner"] for r in records if r.get("owner")}),
            "blocking_issue_count": counts["blocking"],
            "review_issue_count": counts["review"],
            "skipped_sheet_count": len(skipped),
        },
        "skipped_sheets": skipped,
        "issues": issues,
        "records": records,
    }
    (output_dir / "normalized.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    columns = ["source_record_id", *CANONICAL_FIELDS, "source_file", "source_sheet", "source_range", "source_row"]
    with (output_dir / "normalized.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)

    lines = [
        "# OGSM 导入预检", "", f"- 源文件：`{path.name}`", f"- 范围：`{scope}`",
        f"- 行动数：{len(records)}", f"- 策略数：{payload['summary']['strategy_count']}",
        f"- 负责人原文数：{payload['summary']['owner_text_count']}", f"- 阻断问题：{counts['blocking']}",
        f"- 待确认问题：{counts['review']}", "", "## 问题", "",
    ]
    if not issues:
        lines.append("没有发现结构性问题；人员身份、业务阈值和提醒对象仍需在 Base 写入前确认。")
    else:
        for issue in issues:
            lines.append(f"- **{issue['severity']} / {issue['code']}**：{issue['message']}")
    lines.extend(["", "## 行动映射", "", "| 源记录ID | S｜策略 | M｜衡量 | 行动 | 负责人原文 | 截止日期 |", "|---|---|---|---|---|---|"])
    for record in records:
        def safe(value: Any) -> str:
            return str(value or "").replace("|", "\\|").replace("\n", "<br>")
        lines.append("| " + " | ".join(safe(record.get(key)) for key in ("source_record_id", "strategy", "measure", "action", "owner", "due_date")) + " |")
    (output_dir / "precheck.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--scope", choices=("department", "personal"), default="department")
    parser.add_argument("--header-row", type=int, help="1-based header row to use for every sheet")
    args = parser.parse_args()
    if not args.source.exists():
        parser.error(f"source does not exist: {args.source}")
    try:
        records, issues, skipped = parse_source(args.source, args.header_row)
        validate(records, issues, args.scope)
        write_outputs(args.source, args.output_dir, args.scope, records, issues, skipped)
    except Exception as exc:
        print(f"precheck failed: {exc}", file=sys.stderr)
        return 1
    blocking = sum(1 for issue in issues if issue["severity"] == "blocking")
    print(json.dumps({"ok": blocking == 0, "records": len(records), "blocking": blocking, "review": sum(1 for i in issues if i["severity"] == "review"), "output_dir": str(args.output_dir)}, ensure_ascii=False))
    return 2 if blocking else 0


if __name__ == "__main__":
    raise SystemExit(main())
